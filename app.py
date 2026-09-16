from io import BytesIO
import html
import re
import unicodedata

import streamlit as st
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

st.set_page_config(
    page_title="Asistente de Caracterización de Procesos",
    page_icon="📘",
    layout="wide",
)


# Los once campos tienen inicialmente el mismo peso.
CAMPOS = {
    "nombre_proceso": "Nombre del proceso",
    "objetivo_proceso": "Objetivo del proceso",
    "responsable_proceso": "Responsable del proceso",
    "proveedor": "Proveedor",
    "entrada": "Entrada",
    "actividades_principales": "Actividades principales",
    "salida": "Salida",
    "cliente_usuario": "Cliente o usuario",
    "criterio_aceptacion": "Criterio de aceptación de la salida",
    "indicador_proceso": "Indicador del proceso",
    "riesgos_observaciones": "Riesgos u observaciones",
}


# ============================================================
# ESTILOS VISUALES
# ============================================================

st.markdown(
    """
    <style>
        .stApp {
            background-color: #f5f7fa;
        }

        .block-container {
            max-width: 1200px;
            padding-top: 2rem;
            padding-bottom: 3rem;
        }

        .hero {
            background: linear-gradient(135deg, #0b3d70 0%, #155a96 100%);
            color: white;
            padding: 1.6rem 1.8rem;
            border-radius: 16px;
            margin-bottom: 1.2rem;
            box-shadow: 0 8px 24px rgba(11, 61, 112, 0.16);
        }

        .hero h1 {
            margin: 0;
            color: white;
            font-size: clamp(1.8rem, 4vw, 2.7rem);
        }

        .hero p {
            margin: 0.55rem 0 0 0;
            color: #eaf3fb;
            font-size: 1.05rem;
        }

        .section-title {
            color: #0b3d70;
            font-size: 1.35rem;
            font-weight: 700;
            margin-top: 1.2rem;
            margin-bottom: 0.65rem;
        }

        .status-card {
            padding: 1rem 1.2rem;
            border-radius: 12px;
            color: white;
            font-weight: 700;
            font-size: 1.05rem;
            margin: 0.8rem 0 1rem 0;
        }

        .sipoc-card {
            background: white;
            border: 1px solid #dbe3eb;
            border-top: 5px solid #155a96;
            border-radius: 12px;
            padding: 1rem;
            min-height: 190px;
            box-shadow: 0 4px 14px rgba(26, 54, 80, 0.07);
            overflow-wrap: anywhere;
        }

        .sipoc-title {
            color: #0b3d70;
            font-weight: 800;
            font-size: 0.88rem;
            letter-spacing: 0.04rem;
            margin-bottom: 0.6rem;
        }

        .sipoc-content {
            color: #263746;
            font-size: 0.95rem;
            line-height: 1.45;
        }

        .sipoc-arrow {
            text-align: center;
            color: #155a96;
            font-size: 1.4rem;
            font-weight: 700;
            margin-top: 0.3rem;
        }

        .detail-card {
            background: white;
            border: 1px solid #dbe3eb;
            border-radius: 12px;
            padding: 1rem 1.1rem;
            min-height: 135px;
            margin-bottom: 0.8rem;
            overflow-wrap: anywhere;
        }

        .detail-label {
            color: #0b3d70;
            font-weight: 750;
            font-size: 0.9rem;
            margin-bottom: 0.45rem;
        }

        .detail-value {
            color: #263746;
            line-height: 1.5;
        }

        div.stButton > button,
        div.stDownloadButton > button {
            border-radius: 10px;
            font-weight: 700;
            min-height: 3rem;
        }

        div.stDownloadButton > button {
            background: #0b3d70;
            color: white;
            border: 1px solid #0b3d70;
        }

        div.stDownloadButton > button:hover {
            background: #155a96;
            color: white;
            border-color: #155a96;
        }

        @media (max-width: 700px) {
            .block-container {
                padding-left: 1rem;
                padding-right: 1rem;
            }

            .sipoc-card {
                min-height: auto;
            }
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# FUNCIONES DE ESTADO Y VALIDACIÓN
# ============================================================

def inicializar_estado():
    """
    Crea las variables necesarias para conservar el formulario
    y el último análisis durante la sesión.
    """
    for clave in CAMPOS:
        if clave not in st.session_state:
            st.session_state[clave] = ""

    if "ultimo_analisis" not in st.session_state:
        st.session_state.ultimo_analisis = None

    if "mensaje_validacion" not in st.session_state:
        st.session_state.mensaje_validacion = ""


def limpiar_formulario():
    """
    Limpia todos los campos y elimina el último análisis.
    Esta función se ejecuta antes de que Streamlit vuelva
    a dibujar la página.
    """
    for clave in CAMPOS:
        st.session_state[clave] = ""

    st.session_state.ultimo_analisis = None
    st.session_state.mensaje_validacion = ""


def texto_diligenciado(valor):
    """
    Verifica si un campo contiene texto diferente de espacios.
    """
    return bool(str(valor).strip())


def calcular_completitud(datos):
    """
    Calcula el porcentaje según la cantidad de campos diligenciados.
    Todos los campos tienen el mismo peso.
    """
    total_campos = len(CAMPOS)

    campos_diligenciados = sum(
        texto_diligenciado(datos[clave])
        for clave in CAMPOS
    )

    porcentaje = round(
        (campos_diligenciados / total_campos) * 100
    )

    campos_pendientes = [
        etiqueta
        for clave, etiqueta in CAMPOS.items()
        if not texto_diligenciado(datos[clave])
    ]

    return porcentaje, campos_pendientes


def determinar_semaforo(porcentaje):
    """
    Determina el estado, mensaje y color del semáforo.
    """
    if porcentaje >= 80:
        return (
            "VERDE",
            "Caracterización completa.",
            "#218739",
        )

    if porcentaje >= 50:
        return (
            "AMARILLO",
            "Caracterización parcialmente completa.",
            "#c28a00",
        )

    return (
        "ROJO",
        "Caracterización incompleta.",
        "#b42318",
    )


def valor_visible(valor):
    """
    Devuelve 'No diligenciado' cuando el campo está vacío.
    """
    valor = str(valor).strip()
    return valor if valor else "No diligenciado"


def texto_html(valor):
    """
    Prepara el contenido para mostrarlo de forma segura
    dentro de las tarjetas HTML.
    """
    contenido = html.escape(valor_visible(valor))
    return contenido.replace("\n", "<br>")


# ============================================================
# FUNCIONES DE VISUALIZACIÓN
# ============================================================

def mostrar_tarjeta(etiqueta, valor):
    """
    Muestra un dato de la ficha dentro de una tarjeta.
    """
    etiqueta_segura = html.escape(etiqueta)
    valor_seguro = texto_html(valor)

    st.markdown(
        f"""
        <div class="detail-card">
            <div class="detail-label">{etiqueta_segura}</div>
            <div class="detail-value">{valor_seguro}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def mostrar_sipoc(datos):
    """
    Construye una visualización SIPOC sencilla y adaptable.
    """
    nombre_proceso = texto_html(
        datos["nombre_proceso"]
    )

    actividades = texto_html(
        datos["actividades_principales"]
    )

    contenido_proceso = (
        f"<strong>{nombre_proceso}</strong>"
        f"<br><br>{actividades}"
    )

    elementos_sipoc = [
        (
            "PROVEEDOR",
            texto_html(datos["proveedor"]),
        ),
        (
            "ENTRADA",
            texto_html(datos["entrada"]),
        ),
        (
            "PROCESO",
            contenido_proceso,
        ),
        (
            "SALIDA",
            texto_html(datos["salida"]),
        ),
        (
            "CLIENTE",
            texto_html(datos["cliente_usuario"]),
        ),
    ]

    columnas = st.columns(5, gap="small")

    for indice, ((titulo, contenido), columna) in enumerate(
        zip(elementos_sipoc, columnas)
    ):
        with columna:
            st.markdown(
                f"""
                <div class="sipoc-card">
                    <div class="sipoc-title">{titulo}</div>
                    <div class="sipoc-content">{contenido}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )

            if indice < len(elementos_sipoc) - 1:
                st.markdown(
                    '<div class="sipoc-arrow">→</div>',
                    unsafe_allow_html=True,
                )


# ============================================================
# FUNCIONES PARA GENERAR EL DOCUMENTO WORD
# ============================================================

def aplicar_estilo_documento(documento):
    """
    Aplica márgenes, tipografía y colores al documento Word.
    """
    seccion = documento.sections[0]

    seccion.top_margin = Inches(0.7)
    seccion.bottom_margin = Inches(0.7)
    seccion.left_margin = Inches(0.75)
    seccion.right_margin = Inches(0.75)

    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Aptos"
    estilo_normal.font.size = Pt(10.5)

    estilos_titulos = [
        "Title",
        "Heading 1",
        "Heading 2",
    ]

    for nombre_estilo in estilos_titulos:
        estilo = documento.styles[nombre_estilo]
        estilo.font.name = "Aptos Display"
        estilo.font.color.rgb = RGBColor(
            11,
            61,
            112,
        )


def agregar_fila_informacion(tabla, etiqueta, valor):
    """
    Agrega una fila con etiqueta y contenido a una tabla Word.
    """
    celdas = tabla.add_row().cells

    celdas[0].text = etiqueta
    celdas[1].text = valor_visible(valor)

    if celdas[0].paragraphs[0].runs:
        celdas[0].paragraphs[0].runs[0].bold = True

    for celda in celdas:
        celda.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )


def generar_documento_word(analisis):
    """
    Genera el documento Word completamente en memoria.
    Retorna el contenido del archivo como bytes.
    """
    datos = analisis["datos"]

    documento = Document()
    aplicar_estilo_documento(documento)

    # Título principal
    titulo = documento.add_heading(
        "Ficha de Caracterización del Proceso",
        level=0,
    )
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Información general
    documento.add_heading(
        "1. Información general",
        level=1,
    )

    tabla_general = documento.add_table(
        rows=0,
        cols=2,
    )
    tabla_general.style = "Table Grid"

    agregar_fila_informacion(
        tabla_general,
        "Nombre del proceso",
        datos["nombre_proceso"],
    )

    agregar_fila_informacion(
        tabla_general,
        "Objetivo del proceso",
        datos["objetivo_proceso"],
    )

    agregar_fila_informacion(
        tabla_general,
        "Responsable del proceso",
        datos["responsable_proceso"],
    )

    # Tabla SIPOC
    documento.add_paragraph()

    documento.add_heading(
        "2. Tabla SIPOC",
        level=1,
    )

    tabla_sipoc = documento.add_table(
        rows=1,
        cols=5,
    )
    tabla_sipoc.style = "Table Grid"

    encabezados_sipoc = [
        "Proveedor",
        "Entrada",
        "Proceso",
        "Salida",
        "Cliente o usuario",
    ]

    for celda, encabezado in zip(
        tabla_sipoc.rows[0].cells,
        encabezados_sipoc,
    ):
        celda.text = encabezado

        if celda.paragraphs[0].runs:
            celda.paragraphs[0].runs[0].bold = True

        celda.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        celda.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

    fila_sipoc = tabla_sipoc.add_row().cells

    texto_proceso = (
        f"Nombre del proceso:\n"
        f"{valor_visible(datos['nombre_proceso'])}\n\n"
        f"Actividades principales:\n"
        f"{valor_visible(datos['actividades_principales'])}"
    )

    valores_sipoc = [
        valor_visible(datos["proveedor"]),
        valor_visible(datos["entrada"]),
        texto_proceso,
        valor_visible(datos["salida"]),
        valor_visible(datos["cliente_usuario"]),
    ]

    for celda, valor in zip(
        fila_sipoc,
        valores_sipoc,
    ):
        celda.text = valor
        celda.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

    # Información complementaria
    documento.add_paragraph()

    documento.add_heading(
        "3. Información complementaria",
        level=1,
    )

    tabla_complementaria = documento.add_table(
        rows=0,
        cols=2,
    )
    tabla_complementaria.style = "Table Grid"

    agregar_fila_informacion(
        tabla_complementaria,
        "Criterio de aceptación de la salida",
        datos["criterio_aceptacion"],
    )

    agregar_fila_informacion(
        tabla_complementaria,
        "Indicador del proceso",
        datos["indicador_proceso"],
    )

    agregar_fila_informacion(
        tabla_complementaria,
        "Riesgos u observaciones",
        datos["riesgos_observaciones"],
    )

    # Resultado del análisis
    documento.add_paragraph()

    documento.add_heading(
        "4. Resultado del análisis",
        level=1,
    )

    parrafo_porcentaje = documento.add_paragraph()
    etiqueta_porcentaje = parrafo_porcentaje.add_run(
        "Porcentaje de completitud: "
    )
    etiqueta_porcentaje.bold = True

    parrafo_porcentaje.add_run(
        f"{analisis['porcentaje']} %"
    )

    parrafo_estado = documento.add_paragraph()
    etiqueta_estado = parrafo_estado.add_run(
        "Estado del semáforo: "
    )
    etiqueta_estado.bold = True

    parrafo_estado.add_run(
        f"{analisis['estado']} | "
        f"{analisis['mensaje_estado']}"
    )

    documento.add_paragraph(
        "El porcentaje refleja el nivel de diligenciamiento "
        "de la ficha, pero no evalúa por sí solo la calidad "
        "técnica de la información registrada."
    )

    documento.add_heading(
        "Campos pendientes por diligenciar",
        level=2,
    )

    if analisis["pendientes"]:
        for campo in analisis["pendientes"]:
            documento.add_paragraph(
                campo,
                style="List Bullet",
            )
    else:
        documento.add_paragraph(
            "Todos los campos fueron diligenciados."
        )

    # Creación del archivo en memoria
    archivo_memoria = BytesIO()
    documento.save(archivo_memoria)
    archivo_memoria.seek(0)

    return archivo_memoria.getvalue()


def limpiar_nombre_archivo(nombre_proceso):
    """
    Genera un nombre de archivo seguro.

    Elimina tildes, espacios, símbolos y caracteres
    que podrían ocasionar problemas en el nombre del archivo.
    """
    nombre = str(nombre_proceso).strip()

    if not nombre:
        return "Caracterizacion_del_Proceso.docx"

    # Eliminar tildes y otros signos diacríticos
    nombre = unicodedata.normalize(
        "NFKD",
        nombre,
    )

    nombre = "".join(
        caracter
        for caracter in nombre
        if not unicodedata.combining(caracter)
    )

    # Sustituir grupos de caracteres no permitidos
    nombre = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        nombre,
    )

    # Evitar varios guiones bajos seguidos
    nombre = re.sub(
        r"_+",
        "_",
        nombre,
    )

    nombre = nombre.strip("_.-")

    if not nombre:
        nombre = "del_Proceso"

    # Limitar la longitud para evitar nombres excesivamente largos
    nombre = nombre[:100]

    return f"Caracterizacion_{nombre}.docx"


# ============================================================
# INICIO DE LA APLICACIÓN
# ============================================================

inicializar_estado()

st.markdown(
    """
    <div class="hero">
        <h1>Asistente de Caracterización de Procesos</h1>
        <p>
            Construya, revise y descargue de manera sencilla
            la caracterización de su proceso.
        </p>
    </div>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="section-title">'
    '1. Diligencie la información del proceso'
    '</div>',
    unsafe_allow_html=True,
)

st.caption(
    "Complete los campos que conozca. Puede dejar campos "
    "vacíos y volver a analizar más adelante."
)


# ============================================================
# FORMULARIO
# ============================================================

with st.form("formulario_caracterizacion"):

    columna_izquierda, columna_derecha = st.columns(
        2,
        gap="large",
    )

    with columna_izquierda:

        st.text_input(
            "Nombre del proceso",
            key="nombre_proceso",
            placeholder="Ejemplo: Gestión de matrículas",
        )

        st.text_area(
            "Objetivo del proceso",
            key="objetivo_proceso",
            placeholder=(
                "Ejemplo: Garantizar la matrícula oportuna "
                "y correcta de los estudiantes."
            ),
            height=110,
        )

        st.text_input(
            "Responsable del proceso",
            key="responsable_proceso",
            placeholder=(
                "Ejemplo: Dirección de Registro Académico"
            ),
        )

        st.text_area(
            "Proveedor",
            key="proveedor",
            placeholder=(
                "Ejemplo: Aspirante, Facultad, Tesorería"
            ),
            height=100,
        )

        st.text_area(
            "Entrada",
            key="entrada",
            placeholder=(
                "Ejemplo: Documentos de admisión, recibo "
                "de pago y solicitud de matrícula"
            ),
            height=120,
        )

        st.text_area(
            "Actividades principales",
            key="actividades_principales",
            placeholder=(
                "Escriba una actividad por línea.\n"
                "Ejemplo:\n"
                "Validar documentos\n"
                "Verificar pago\n"
                "Registrar asignaturas"
            ),
            height=180,
        )

    with columna_derecha:

        st.text_area(
            "Salida",
            key="salida",
            placeholder=(
                "Ejemplo: Estudiante matriculado y "
                "horario generado"
            ),
            height=110,
        )

        st.text_area(
            "Cliente o usuario",
            key="cliente_usuario",
            placeholder=(
                "Ejemplo: Estudiantes, facultades y "
                "dependencias académicas"
            ),
            height=100,
        )

        st.text_area(
            "Criterio de aceptación de la salida",
            key="criterio_aceptacion",
            placeholder=(
                "Ejemplo: Matrícula activa, datos completos "
                "y horario confirmado"
            ),
            height=110,
        )

        st.text_area(
            "Indicador del proceso",
            key="indicador_proceso",
            placeholder=(
                "Ejemplo: Porcentaje de matrículas "
                "completadas dentro del plazo"
            ),
            height=110,
        )

        st.text_area(
            "Riesgos u observaciones",
            key="riesgos_observaciones",
            placeholder=(
                "Ejemplo: Pagos no aplicados, documentos "
                "incompletos o fallas de plataforma"
            ),
            height=150,
        )

    boton_analizar = st.form_submit_button(
        "Analizar proceso",
        type="primary",
        use_container_width=True,
    )


# El botón se encuentra fuera del formulario para que pueda
# ejecutar su función de limpieza de forma independiente.
st.button(
    "Limpiar formulario",
    on_click=limpiar_formulario,
    use_container_width=True,
)


# ============================================================
# EJECUCIÓN DEL ANÁLISIS
# ============================================================

if boton_analizar:

    datos_actuales = {
        clave: st.session_state.get(clave, "")
        for clave in CAMPOS
    }

    existe_informacion = any(
        texto_diligenciado(valor)
        for valor in datos_actuales.values()
    )

    if not existe_informacion:

        st.session_state.ultimo_analisis = None

        st.session_state.mensaje_validacion = (
            "Debe diligenciar al menos algunos campos "
            "antes de analizar el proceso."
        )

    else:

        porcentaje, pendientes = calcular_completitud(
            datos_actuales
        )

        estado, mensaje_estado, color = (
            determinar_semaforo(porcentaje)
        )

        # Se guarda una copia de los datos.
        # Los cambios posteriores en el formulario no alteran
        # el Word hasta que se vuelva a analizar.
        st.session_state.ultimo_analisis = {
            "datos": datos_actuales.copy(),
            "porcentaje": porcentaje,
            "pendientes": pendientes,
            "estado": estado,
            "mensaje_estado": mensaje_estado,
            "color": color,
        }

        st.session_state.mensaje_validacion = ""


if st.session_state.mensaje_validacion:
    st.warning(
        st.session_state.mensaje_validacion
    )


# ============================================================
# PRESENTACIÓN DEL ÚLTIMO ANÁLISIS
# ============================================================

analisis = st.session_state.ultimo_analisis

if analisis:

    st.divider()

    st.markdown(
        '<div class="section-title">'
        '2. Resultado del análisis'
        '</div>',
        unsafe_allow_html=True,
    )

    columna_metrica, columna_progreso = st.columns(
        [1, 2],
        gap="large",
    )

    with columna_metrica:
        st.metric(
            "Completitud de la caracterización",
            f"{analisis['porcentaje']} %",
        )

    with columna_progreso:
        st.write("Progreso de diligenciamiento")

        st.progress(
            analisis["porcentaje"] / 100
        )

    st.info(
        "El porcentaje refleja el nivel de diligenciamiento "
        "de la ficha, pero no evalúa por sí solo la calidad "
        "técnica de la información registrada."
    )

    st.markdown(
        f"""
        <div
            class="status-card"
            style="background-color: {analisis['color']};"
        >
            {analisis['estado']} |
            {analisis['mensaje_estado']}
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        "#### Campos pendientes por diligenciar"
    )

    if analisis["pendientes"]:
        for campo in analisis["pendientes"]:
            st.markdown(f"- {campo}")
    else:
        st.success(
            "Todos los campos fueron diligenciados."
        )

    # Visualización SIPOC
    st.markdown(
        '<div class="section-title">'
        '3. Visualización SIPOC'
        '</div>',
        unsafe_allow_html=True,
    )

    mostrar_sipoc(
        analisis["datos"]
    )

    # Resumen completo
    st.markdown(
        '<div class="section-title">'
        '4. Resumen de la caracterización'
        '</div>',
        unsafe_allow_html=True,
    )

    datos_analizados = analisis["datos"]
    elementos_resumen = list(CAMPOS.items())

    columna_resumen_1, columna_resumen_2 = st.columns(
        2,
        gap="large",
    )

    with columna_resumen_1:
        for clave, etiqueta in elementos_resumen[:6]:
            mostrar_tarjeta(
                etiqueta,
                datos_analizados[clave],
            )

    with columna_resumen_2:
        for clave, etiqueta in elementos_resumen[6:]:
            mostrar_tarjeta(
                etiqueta,
                datos_analizados[clave],
            )

    columna_analisis_1, columna_analisis_2 = st.columns(
        2,
        gap="large",
    )

    with columna_analisis_1:

        mostrar_tarjeta(
            "Porcentaje de completitud",
            f"{analisis['porcentaje']} %",
        )

        mostrar_tarjeta(
            "Estado del semáforo",
            (
                f"{analisis['estado']} | "
                f"{analisis['mensaje_estado']}"
            ),
        )

    with columna_analisis_2:

        if analisis["pendientes"]:
            texto_pendientes = "\n".join(
                f"• {campo}"
                for campo in analisis["pendientes"]
            )
        else:
            texto_pendientes = (
                "Todos los campos fueron diligenciados."
            )

        mostrar_tarjeta(
            "Campos pendientes por diligenciar",
            texto_pendientes,
        )

    # Descarga del Word
    st.markdown(
        '<div class="section-title">'
        '5. Descargar documento'
        '</div>',
        unsafe_allow_html=True,
    )

    archivo_word = generar_documento_word(
        analisis
    )

    nombre_archivo = limpiar_nombre_archivo(
        datos_analizados["nombre_proceso"]
    )

    st.download_button(
        label="Descargar caracterización en Word",
        data=archivo_word,
        file_name=nombre_archivo,
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        use_container_width=True,
    )
